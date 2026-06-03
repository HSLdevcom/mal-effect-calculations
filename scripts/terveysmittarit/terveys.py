import os.path
import pandas
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows


# RESULT_DIR = "Helmet-ja-mittariajot\\Malliajojen_tulokset_suunnitelma_kesa2023"
RESULT_DIR = "T:\\JohannaP\\mal2023"
SCENARIO_PATH = "terveysmittariherkkyystarkastelut\\terveysvaikutusmittarit{}.xlsx"


def write_table(worksheet, scenario, file_name, row_nr):
    df = pandas.read_csv(
        os.path.join(RESULT_DIR, scenario, file_name), delim_whitespace=True)
    rows = dataframe_to_rows(df, index=True, header=False)
    for r_idx, row in enumerate(rows, row_nr):
        for c_idx, value in enumerate(row, 1):
            worksheet.cell(row=r_idx, column=c_idx, value=value)


def write_worksheet(workbook, alternative, scenario):
    worksheet = workbook["INPUT_H41_{}".format(alternative)]
    scenario = alternative + scenario
    write_table(worksheet, scenario, "vehicle_kms_vdfs.txt", 2)
    write_table(worksheet, scenario, "vehicle_kms_areas.txt", 12)
    write_table(worksheet, scenario, "vehicle_kms_vdfs_areas.txt", 21)
    write_table(worksheet, scenario, "trips_areas.txt", 41)
    write_table(worksheet, scenario, "noise_areas.txt", 53)


def write_workbook(workbook_path, scenario="", only_ve0=False):
    workbook_path = workbook_path.format(scenario)
    wb = load_workbook(workbook_path)
    write_worksheet(wb, "2040_ve0", scenario)
    if not only_ve0:
        # write_worksheet(wb, "2040_ve1", scenario)
        write_worksheet(wb, "2040_suunnitelma", scenario)
    wb.save(workbook_path)


write_workbook("MAL 2023 terveysvaikutusmittarit.xlsx")
write_workbook(SCENARIO_PATH, "_etatyo")
write_workbook(SCENARIO_PATH, "_jklmaksi")
write_workbook(SCENARIO_PATH, "_mini")
write_workbook(SCENARIO_PATH, "_tieruuhka")
write_workbook(SCENARIO_PATH, "_ruuhkautuva", only_ve0=True)
workbook_path = "MAL 2023 terveysvaikutusmittarit.xlsx"
wb = load_workbook(workbook_path)
write_worksheet(wb, "2018", "")
wb.save(workbook_path)
